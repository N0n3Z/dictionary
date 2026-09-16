#!/usr/bin/env python3
"""
Écrit le classeur "vue historique par variable sémantique" à partir du JSON produit
par 05_build_historique.py : une ligne par génération de Code_IPCAL (une seule pour
l'immense majorité des codes, plusieurs pour les codes ayant subi une rupture
sémantique suspectée), avec un libellé par année en colonnes pour un scan visuel
rapide, plus une feuille dédiée listant sans ambiguïté les codes scindés (celle à
regarder avant de construire un mapping vers des agrégats nationaux), et le détail
de chaque coupure sémantique.

Usage :
    python3 06_write_historique_excel.py data/historique.json -o IPCAL_historique_variables.xlsx
"""
import json, argparse, os, re as _re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ILLEGAL = _re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

def san(v):
    return ILLEGAL.sub('', v) if isinstance(v, str) else v


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('historique', help='JSON produit par 05_build_historique.py')
    ap.add_argument('-o', '--out', required=True)
    args = ap.parse_args()

    data = json.load(open(args.historique, encoding='utf-8'))
    years = data['annees']
    variables = data['variables']
    ruptures = data['ruptures_semantiques']
    codes_scindes = sorted({v['Code_IPCAL'] for v in variables if v['Rupture_semantique']})
    print(f'{len(variables)} lignes (générations), {len(codes_scindes)} codes scindés, {len(years)} années ({years[0]}-{years[-1]})')

    wb = openpyxl.Workbook()
    HEAD = PatternFill('solid', fgColor='1F4E78')
    HEADF = Font(name='Arial', bold=True, color='FFFFFF', size=9)
    CELLF = Font(name='Arial', size=9)
    thin = Side(style='thin', color='D9D9D9')
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt = PatternFill('solid', fgColor='F2F6FB')
    gap_fill = PatternFill('solid', fgColor='FDE9D9')        # année hors génération / absente
    flag_fill = PatternFill('solid', fgColor='F8696B')       # rupture sémantique -- très visible
    flag_font = Font(name='Arial', bold=True, color='FFFFFF', size=9)

    # ---- Feuille 0 : Codes_avec_rupture (LA feuille à consulter avant tout mapping) ----
    cs = wb.active
    cs.title = 'Codes_avec_rupture'
    cs_cols = ['Code_IPCAL', 'Nb_generations', 'Cle_mapping', 'Validite_debut', 'Validite_fin', 'Libelle']
    cs.append(['⚠ ' + c if c == 'Code_IPCAL' else c for c in cs_cols])
    for c in range(1, len(cs_cols) + 1):
        cell = cs.cell(row=1, column=c)
        cell.fill = flag_fill; cell.font = flag_font
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        cell.border = BORDER
    by_code_gens = {}
    for v in variables:
        if v['Rupture_semantique']:
            by_code_gens.setdefault(v['Code_IPCAL'], []).append(v)
    rr = 2
    for code in codes_scindes:
        gens = sorted(by_code_gens[code], key=lambda v: v['Generation'])
        for v in gens:
            libelle = v['Libelle_par_annee'].get(str(v['Validite_fin']), v['Libelle_par_annee'].get(v['Validite_fin'], ''))
            row = [code, v['Nb_generations_total'], v['Cle_mapping'], v['Validite_debut'], v['Validite_fin'], libelle]
            cs.append([san(x) for x in row])
            for c in range(1, len(cs_cols) + 1):
                cell = cs.cell(row=rr, column=c); cell.border = BORDER
                cell.font = Font(name='Arial', bold=(c == 1), size=9)
                if c == 1:
                    cell.fill = PatternFill('solid', fgColor='FDE9D9')
            rr += 1
    cs.freeze_panes = 'A2'
    cs.auto_filter.ref = f"A1:{get_column_letter(len(cs_cols))}{rr - 1}"
    CSW = {'Code_IPCAL': 12, 'Nb_generations': 14, 'Cle_mapping': 16, 'Validite_debut': 13, 'Validite_fin': 12, 'Libelle': 46}
    for i, k in enumerate(cs_cols, start=1):
        cs.column_dimensions[get_column_letter(i)].width = CSW.get(k, 16)
    cs.row_dimensions[1].height = 30
    if not codes_scindes:
        cs.cell(row=2, column=1, value="Aucun code scindé dans ce jeu de données.").font = Font(name='Arial', italic=True, size=10)

    # ---- Feuille 1 : Historique (une ligne par génération) ----
    ws = wb.create_sheet('Historique')
    fixed_cols = ['Code_IPCAL', 'Rupture_semantique', 'Cle_mapping', 'Generation', 'Nb_generations_total',
                  'Prefixe', 'Conjoint', 'Nature_variable', 'Validite_debut', 'Validite_fin',
                  'Nb_annees_presentes', 'Continuite_pct', 'Rupture_disponibilite_mineure']
    year_cols = [f'Libelle_{y}' for y in years]
    cols = fixed_cols + year_cols

    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD; cell.font = HEADF
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        cell.border = BORDER

    variables_sorted = sorted(variables, key=lambda v: (v['Code_IPCAL'], v['Generation']))
    for i, v in enumerate(variables_sorted):
        row_vals = [v.get(k, '') for k in fixed_cols]
        libelles = v['Libelle_par_annee']
        row_vals += [libelles.get(str(y), libelles.get(y, '')) for y in years]
        ws.append([san(x) for x in row_vals])
        rr2 = i + 2
        gen_years = set(v['Annees_presentes'])
        span = range(v['Validite_debut'], v['Validite_fin'] + 1)
        for c, k in enumerate(cols, start=1):
            cell = ws.cell(row=rr2, column=c); cell.font = CELLF; cell.border = BORDER
            if k == 'Rupture_semantique':
                cell.alignment = Alignment(horizontal='center')
                if v['Rupture_semantique']:
                    cell.fill = flag_fill; cell.font = flag_font
                    cell.value = '⚠ OUI'
                else:
                    cell.value = 'Non'
            elif k == 'Continuite_pct':
                cell.alignment = Alignment(horizontal='right')
            elif k.startswith('Libelle_'):
                y = int(k.split('_')[1])
                if y not in range(v['Validite_debut'], v['Validite_fin'] + 1):
                    cell.fill = PatternFill('solid', fgColor='E7E6E6')  # hors génération (autre plage / pas encore/plus présent)
                elif y not in gen_years:
                    cell.fill = gap_fill  # coupure mineure interne à la génération
                elif i % 2 == 1:
                    cell.fill = alt
            elif i % 2 == 1 and k != 'Rupture_semantique':
                cell.fill = alt
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(variables_sorted) + 1}"
    W = {'Nature_variable': 32, 'Rupture_semantique': 13, 'Cle_mapping': 15, 'Continuite_pct': 11,
         'Rupture_disponibilite_mineure': 14, 'Nb_generations_total': 12}
    for i, k in enumerate(cols, start=1):
        width = W.get(k, 30 if k.startswith('Libelle_') else 13)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 40

    # ---- Feuille 2 : Ruptures_semantiques (une ligne par coupure = frontière entre générations) ----
    al = wb.create_sheet('Ruptures_semantiques')
    al_cols = ['Code_IPCAL', 'Nature_variable', 'Cle_mapping_avant', 'Cle_mapping_apres', 'annees_absentes',
               'avant_annee', 'avant_libelle', 'avant_chemin', 'apres_annee', 'apres_libelle', 'apres_chemin', 'similarite']
    al_headers = ['Code_IPCAL', 'Nature_variable', 'Cle_mapping (avant)', 'Cle_mapping (après)', 'Années absentes',
                  'Avant : année', 'Avant : libellé', 'Avant : hiérarchie', 'Après : année', 'Après : libellé',
                  'Après : hiérarchie', 'Similarité']
    al.append(al_headers)
    for c in range(1, len(al_cols) + 1):
        cell = al.cell(row=1, column=c)
        cell.fill = HEAD; cell.font = HEADF
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        cell.border = BORDER

    ruptures_sorted = sorted(ruptures, key=lambda a: a['similarite'])  # cas les plus suspects d'abord
    for i, a in enumerate(ruptures_sorted):
        row = [a['Code_IPCAL'], a['Nature_variable'], a['Cle_mapping_avant'], a['Cle_mapping_apres'],
               ','.join(str(y) for y in a['annees_absentes']), a['avant_annee'], a['avant_libelle'],
               a['avant_chemin'], a['apres_annee'], a['apres_libelle'], a['apres_chemin'], a['similarite']]
        al.append([san(x) for x in row])
        rr3 = i + 2
        for c in range(1, len(al_cols) + 1):
            cell = al.cell(row=rr3, column=c); cell.font = CELLF; cell.border = BORDER
            if i % 2 == 1:
                cell.fill = alt
    if ruptures_sorted:
        al.freeze_panes = 'E2'
        al.auto_filter.ref = f"A1:{get_column_letter(len(al_cols))}{len(ruptures_sorted) + 1}"
    ALW = {'avant_libelle': 36, 'apres_libelle': 36, 'avant_chemin': 50, 'apres_chemin': 50,
           'Nature_variable': 32, 'annees_absentes': 16, 'Cle_mapping_avant': 15, 'Cle_mapping_apres': 15}
    for i, k in enumerate(al_cols, start=1):
        al.column_dimensions[get_column_letter(i)].width = ALW.get(k, 13)
    al.row_dimensions[1].height = 30

    # ---- Feuille 3 : Légende / méthode ----
    lg = wb.create_sheet('Legende')
    lg.column_dimensions['A'].width = 34; lg.column_dimensions['B'].width = 112
    tf = Font(name='Arial', bold=True, size=13, color='1F4E78')
    sf = Font(name='Arial', bold=True, size=11, color='1F4E78')
    bf = Font(name='Arial', bold=True, size=10); nf = Font(name='Arial', size=10)
    def put(row, a, b):
        ca = lg.cell(row=row, column=1, value=a); ca.font = bf; ca.alignment = Alignment(wrap_text=True, vertical='top')
        cb = lg.cell(row=row, column=2, value=b); cb.font = nf; cb.alignment = Alignment(wrap_text=True, vertical='top')
    r = 1
    lg.cell(row=r, column=1, value='Historique multi-année par variable sémantique — IPCAL').font = tf; r += 2
    lg.cell(row=r, column=1, value='PRINCIPE').font = sf; r += 1
    put(r, 'Pas d\'identifiant inventé', "Code_IPCAL reste la clé partout, y compris dans le Dictionnaire annuel (Code_IPCAL + Annee_revenus désambiguïse déjà totalement). Ce classeur ne scinde en plusieurs lignes ('générations') que les codes ayant une rupture sémantique suspectée -- utile uniquement pour un usage hors-année, typiquement un mapping vers des agrégats nationaux."); r += 1
    put(r, 'Rupture_semantique (⚠)', "True sur TOUTES les générations d'un Code_IPCAL scindé. Avant de mapper ce code vers un agrégat national, consulter l'onglet Codes_avec_rupture : autant de lignes que de générations, chacune avec sa plage de validité -- ne jamais mapper Code_IPCAL seul dans ce cas, toujours (Code_IPCAL, année) ou Cle_mapping."); r += 1
    put(r, 'Cle_mapping', "Code_IPCAL seul si une seule génération (~99% des codes -- coût nul). Sinon \"<Code_IPCAL>-<Validite_debut>\" (ex. \"A0270-2021\"), lisible sans table de correspondance. Recommandé comme clé de la future table de mapping ; pour joindre aux données annuelles : Code_IPCAL égal ET année comprise entre Validite_debut et Validite_fin."); r += 1
    put(r, 'Validite_debut / Validite_fin', "Première/dernière année connue de CETTE génération (pas du code entier si celui-ci est scindé)."); r += 1
    put(r, 'Continuite_pct', "Part des années de [Validite_debut, Validite_fin] où le code est effectivement présent (100% = aucune coupure mineure)."); r += 1
    put(r, 'Rupture_disponibilite_mineure', "Coupure de disponibilité interne à cette génération, mais au libellé jugé stable (similarité >= 0.6) -- pas traitée comme un changement de sens, donc pas de scission."); r += 1
    put(r, 'Cellules grises (Historique)', "Année hors de la plage de validité de cette génération (appartient à une autre génération du même code)."); r += 1
    put(r, 'Cellules oranges (Historique)', "Année dans la plage de validité mais code absent (coupure mineure)."); r += 1
    lg.cell(row=r, column=1, value='SEUIL DE DÉTECTION').font = sf; r += 1
    put(r, 'Rupture_semantique (par coupure)', "Similarité difflib du libellé+hiérarchie avant/après < 0.6, après normalisation (minuscules, ponctuation et années à 4 chiffres retirées). Seuil calibré sur 2017-2023 : parmi les 76 coupures observées, la similarité max était 0.559 et correspondait systématiquement, à relecture manuelle, à un vrai changement de sens (ex. A7990 'Option sur action' -> 'Periode 2 : juin - sept'). Ce constat pourrait ne pas se généraliser à des données futures : signal de priorisation, pas de certitude absolue -- vérifier Avant/Après dans Ruptures_semantiques avant de conclure."); r += 1
    lg.cell(row=r, column=1, value='LIMITES CONNUES').font = sf; r += 1
    put(r, 'Fenêtre observée', "'Validite_debut/fin' sont relatives aux données chargées, pas à l'existence réelle du code avant/après cette fenêtre."); r += 1
    put(r, 'Faux négatifs possibles', "Un changement de sens SANS coupure de disponibilité (code présent en continu mais dont la signification évolue) n'est pas détecté -- non observé sur 2017-2023, mais pas prouvé impossible. Un mapping basé sur Rupture_semantique seul resterait alors faux sans aucun signal d'alerte."); r += 1
    put(r, 'Voir aussi', "CLAUDE.md à la racine du dépôt, section 'Dérive sémantique multi-année', et scripts/05_build_historique.py pour le détail de l'algorithme.")

    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or '.', exist_ok=True)
    wb.save(args.out)
    print(f'Écrit: {args.out} | {len(variables)} lignes, {len(codes_scindes)} codes scindés, {len(ruptures)} coupures sémantiques')


if __name__ == '__main__':
    main()
